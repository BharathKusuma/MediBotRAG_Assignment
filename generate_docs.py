import os
import sys
from pathlib import Path
from typing import Any, List, Dict
from docx import Document
from docx.document import Document as DocxDocument
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfgen import canvas

OUTPUT_DIR = Path("e:/Tutorials/CodeBasics/AI_Engineering/12_13Sep/Assignment-3-Evaluation_Guardrail_Pipeline/MediBotRAG_Assignment")

# ==========================================
# 1. GENERATE WORD (.DOCX) DOCUMENT
# ==========================================
def add_custom_heading(doc: DocxDocument, text: str, level: int = 1):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    h.paragraph_format.space_after = Pt(4)
    run = h.add_run(text)
    run.font.name = "Arial"
    run.font.bold = True
    if level == 1:
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(2, 132, 199)
    else:
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(15, 23, 42)
    return h

def create_docx(filename: Path):
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Title & Header
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = title_p.add_run("🏥 MediBot System Architecture & Code Flow")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(15, 23, 42)

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = sub_p.add_run("Enterprise Medical Assistant with Database-Layer RBAC, Hybrid Search & SQL RAG\nMediAssist Health Network")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(11)
    run_sub.font.color.rgb = RGBColor(71, 85, 105)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # 1. Executive Summary
    add_custom_heading(doc, "1. Executive Summary", level=1)
    doc.add_paragraph(
        "MediBot is an enterprise AI assistant designed for MediAssist Health Network (12 hospitals and 40+ clinics). "
        "It addresses two primary operational challenges:\n"
        "• High-precision multimodal knowledge retrieval across 5 distinct domains (clinical, nursing, billing, equipment, and general policies).\n"
        "• Zero-trust Role-Based Access Control (RBAC) enforced directly inside the database retrieval layer, guaranteeing that unauthorized users cannot extract restricted data even through adversarial prompt injection jailbreaks."
    )

    # 2. End-to-End System Flow Diagram
    add_custom_heading(doc, "2. High-Level System Flow Diagram", level=1)

    flowchart_text = (
        "+-----------------------------------------------------------------------------------------+\n"
        "|  Staff User (Doctor / Nurse / Billing / Technician / Admin)                              |\n"
        "|  POST /chat { question, role, auth_token }                                               |\n"
        "+--------------------------------------------+--------------------------------------------+\n"
        "                                             |                                             \n"
        "                                             v                                             \n"
        "                       [ FastAPI Backend Gateway (backend/main.py) ]                      \n"
        "                                             |                                             \n"
        "                        +--------------------+--------------------+                        \n"
        "                        | 1. Security Guardrails & Verification  |                        \n"
        "                        | • Gibberish filter & API-key protection|                        \n"
        "                        +--------------------+--------------------+                        \n"
        "                                             |                                             \n"
        "                                 [ Intent Classifier ]                                     \n"
        "                              /                         \\                                  \n"
        "               (Analytical/Claims Question)         (Policy/Clinical Manual Question)     \n"
        "                            /                             \\                                \n"
        "                           v                               v                               \n"
        "     +--------------------------------+           +--------------------------------+       \n"
        "     | BRANCH A: SQL RAG ENGINE       |           | BRANCH B: HYBRID DOCUMENT RAG  |       \n"
        "     | (backend/rag/sql_rag.py)       |           | (backend/rag/hybrid_retriever) |       \n"
        "     +--------------------------------+           +--------------------------------+       \n"
        "     | 1. RBAC Perms Check            |           | 1. Collection RBAC Pre-Check   |       \n"
        "     |    (Billing Exec / Admin only) |           | 2. Qdrant Vector Filter        |       \n"
        "     | 2. NL-to-SQL Generation (LLM)  |           |    must=[access_roles == role] |       \n"
        "     | 3. Safety Sanitizer (SELECT)   |           | 3. Dual Dense + Sparse Search  |       \n"
        "     | 4. Query SQLite mediassist.db  |           | 4. Reciprocal Rank Fusion (RRF)|       \n"
        "     | 5. LLM Answer Synthesis        |           | 5. Cross-Encoder Rerank (Top 3)|       \n"
        "     +----------------+---------------+           | 6. Grounded LLM Synthesis      |       \n"
        "                      |                           +----------------+---------------+       \n"
        "                      \\                                           /                        \n"
        "                       \\-------------------+---------------------/                         \n"
        "                                           v                                               \n"
        "                       [ Structured JSON Response Returned to UI ]                         \n"
        "                       • Answer text with citations & department badges                    \n"
        "                       • Verified Sources [{ doc, section, score }]                        \n"
        "                       • RBAC security alert status flag                                   \n"
        "+-----------------------------------------------------------------------------------------+"
    )
    p_code = doc.add_paragraph()
    r = p_code.add_run(flowchart_text)
    r.font.name = "Consolas"
    r.font.size = Pt(8.5)
    p_code.paragraph_format.space_after = Pt(12)

    # 3. Phase 1: Ingestion
    add_custom_heading(doc, "3. Phase 1: Ingestion Pipeline (Offline Data Processing)", level=1)
    doc.add_paragraph(
        "Before queries can be answered, raw documents are processed through an automated indexing pipeline:\n"
        "1. Hierarchical Parsing (backend/ingestion/parser.py): Reads PDFs and Markdown files across 5 departments. Preserves Markdown tables, headings, and code blocks. Generates contextual chunks with breadcrumbs like '[Document > Section > Subsection]'.\n"
        "2. Role Tagging: Each chunk is stamped with allowed roles based on folder location (e.g., 'clinical' -> ['doctor', 'admin']; 'billing' -> ['billing_executive', 'admin']).\n"
        "3. Hybrid Vector Embedding (backend/ingestion/embedder.py):\n"
        "   • Dense Vector: 384 dimensions using BAAI/bge-small-en-v1.5 to capture conceptual semantics.\n"
        "   • Sparse BM25 Vector: Generated via FastEmbed Qdrant/bm25 to capture exact medical tokens, ICD-10 codes, drug dosages, and device model numbers.\n"
        "4. Qdrant Storage (backend/ingestion/index_qdrant.py): Stores vectors in collection 'mediassist_knowledge' with dedicated payload indexes on 'access_roles' and 'collection'."
    )

    # 4. Phase 2: Runtime Flow
    add_custom_heading(doc, "4. Phase 2: Runtime Query Processing Flow", level=1)

    add_custom_heading(doc, "Step 1: Request Gateway & Pre-Flight Guardrails", level=2)
    doc.add_paragraph(
        "• Incoming HTTP POST /chat request receives { question, role }.\n"
        "• Gibberish Detector checks vowel-to-consonant ratios and character repetition to intercept nonsensical inputs.\n"
        "• API Key Detection Guardrail blocks suspicious prompt patterns attempting to extract API keys or passwords."
    )

    add_custom_heading(doc, "Step 2: Query Intent Classification", level=2)
    doc.add_paragraph(
        "• The system determines whether the question requires numerical/aggregative SQL analysis or textual knowledge retrieval:\n"
        "  - Analytical Queries: Contain quantitative keywords (count, average, sum, total, approved, claims, tickets, escalated) -> Dispatched to SQL RAG.\n"
        "  - Protocol & Manual Queries: Involve procedures, protocols, drug guidelines, instructions -> Dispatched to Hybrid RAG."
    )

    add_custom_heading(doc, "Step 3A: SQL RAG Engine Execution Path", level=2)
    doc.add_paragraph(
        "1. Role Verification: Only 'billing_executive' and 'admin' are permitted to query financial and maintenance metrics. Non-authorized roles immediately receive an informative refusal message.\n"
        "2. Natural Language to SQL: Gemini/LLM translates the question into valid SQLite SQL referencing tables 'claims' and 'maintenance_tickets'.\n"
        "3. SQL Sanitizer: Strips markdown ticks, strictly verifies the statement begins with SELECT, and blocks data mutation commands (DROP, DELETE, UPDATE, INSERT, ALTER).\n"
        "4. Query Execution: Runs on mediassist.db; results are converted to structured JSON.\n"
        "5. Natural Language Answer Generation: The LLM drafts a concise summary explaining the data findings."
    )

    add_custom_heading(doc, "Step 3B: Hybrid Document RAG Execution Path", level=2)
    doc.add_paragraph(
        "1. Domain Check: Identifies if a role asks about forbidden departments and issues a clear policy explanation.\n"
        "2. Retrieval-Layer RBAC Filtering: Crucially, Qdrant executes filtering at the vector engine level:\n"
        "   filter = models.Filter(must=[FieldCondition(key='access_roles', match=MatchValue(value=role))])\n"
        "   Unauthorized documents physically never reach the server memory, thwarting all prompt jailbreaks.\n"
        "3. Dual Retrieval & RRF Fusion: Simultaneously executes Dense vector search and BM25 sparse search. Results are merged via Reciprocal Rank Fusion (RRF) to retrieve top 10 candidates.\n"
        "4. Cross-Encoder Reranking (backend/rag/reranker.py): Jointly scores the user query with each candidate using 'cross-encoder/ms-marco-MiniLM-L-6-v2', selecting the top 3 most relevant passages.\n"
        "5. LLM Synthesis: Prompts the LLM with strict grounding instructions to synthesize the final answer referencing verified document names and section titles."
    )

    # 5. User Roles & Access Matrix
    add_custom_heading(doc, "5. User Roles & Access Control Matrix", level=1)

    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    headers = ["Role", "Department", "Accessible Collections", "SQL RAG Allowed?", "Demo User"]
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        hdr_cells[i].paragraphs[0].runs[0].font.bold = True
        hdr_cells[i].paragraphs[0].runs[0].font.size = Pt(9.5)

    matrix_data = [
        ("doctor", "Clinical Cardiology", "clinical, nursing, general", "No", "dr.mehta"),
        ("nurse", "ICU & Patient Care", "nursing, general", "No", "nurse.priya"),
        ("billing_executive", "Billing & Insurance", "billing, general", "Yes", "billing.ravi"),
        ("technician", "Biomedical Eng.", "equipment, general", "No", "tech.anand"),
        ("admin", "Executive / IT", "All collections", "Yes", "admin.sys")
    ]

    for row_data in matrix_data:
        row_cells = table.add_row().cells
        for i, val in enumerate(row_data):
            row_cells[i].text = val
            row_cells[i].paragraphs[0].runs[0].font.size = Pt(9)

    # 6. Core Files Table
    add_custom_heading(doc, "6. Key Codebase Modules", level=1)

    file_table = doc.add_table(rows=1, cols=3)
    file_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    f_hdr = file_table.rows[0].cells
    f_hdr[0].text = "Module File"
    f_hdr[1].text = "Layer"
    f_hdr[2].text = "Primary Functionality"
    for cell in f_hdr:
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.size = Pt(9.5)

    files_data = [
        ("backend/main.py", "API Gateway", "FastAPI endpoints (/chat, /login, /collections, /health, /demo-users)."),
        ("backend/ingestion/parser.py", "Data Ingestion", "Parses PDFs and Markdown, builds hierarchical context & chunks."),
        ("backend/ingestion/embedder.py", "Embedding", "Generates dense (bge-small-en) & sparse (BM25) vector representations."),
        ("backend/ingestion/index_qdrant.py", "Vector Storage", "Creates collections, defines schemas, and upserts chunk points."),
        ("backend/rag/hybrid_retriever.py", "Retrieval Core", "Executes Qdrant RBAC filtering, RRF fusion, and pipeline coordination."),
        ("backend/rag/reranker.py", "Reranking", "Cross-Encoder reranking using ms-marco-MiniLM-L-6-v2 down to Top 3."),
        ("backend/rag/sql_rag.py", "SQL Analytics", "NL-to-SQL generation, query sanitizer, SQLite execution on claims/tickets."),
        ("backend/rag/llm_service.py", "LLM Inference", "Multi-model client supporting Gemini, Groq, OpenAI & local fallback.")
    ]

    for f_row in files_data:
        cells = file_table.add_row().cells
        for idx, text in enumerate(f_row):
            cells[idx].text = text
            cells[idx].paragraphs[0].runs[0].font.size = Pt(8.5)

    doc.save(str(filename))
    print(f"Word document saved to: {filename}")

# ==========================================
# 2. GENERATE PDF DOCUMENT (REPORTLAB)
# ==========================================
class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args: Any, **kwargs: Any) -> None:
        super().__init__(*args, **kwargs)
        self._saved_page_states: List[Dict[str, Any]] = []

    def showPage(self) -> None:
        self._saved_page_states.append(dict(self.__dict__))
        # Call protected method via getattr to avoid static typing issues
        start_page_fn = getattr(super(), "_startPage", None) or getattr(self, "_startPage", None)
        if callable(start_page_fn):
            start_page_fn()

    def save(self) -> None:
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_number(num_pages)
            super().showPage()
        super().save()

    def draw_page_number(self, page_count: int) -> None:
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#64748b"))
        # Header rule
        self.setStrokeColor(colors.HexColor("#e2e8f0"))
        self.setLineWidth(0.5)
        self.line(40, letter[1] - 35, letter[0] - 40, letter[1] - 35)
        self.drawString(40, letter[1] - 30, "MediBot System Architecture & Code Flow — MediAssist Health Network")

        # Footer
        page_num = getattr(self, "_pageNumber", 1)
        page_text = f"Page {page_num} of {page_count}"
        self.drawRightString(letter[0] - 40, 25, page_text)
        self.drawString(40, 25, "Confidential — Internal AI Engineering Documentation")
        self.restoreState()

def create_pdf(filename: Path):
    doc = SimpleDocTemplate(
        str(filename),
        pagesize=letter,
        leftMargin=40,
        rightMargin=40,
        topMargin=45,
        bottomMargin=40
    )

    styles = getSampleStyleSheet()

    # Custom styles
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=20,
        leading=24,
        textColor=colors.HexColor("#0f172a"),
        alignment=1,
        spaceAfter=4
    )
    subtitle_style = ParagraphStyle(
        'DocSubTitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#475569"),
        alignment=1,
        spaceAfter=15
    )
    h1_style = ParagraphStyle(
        'H1',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=13,
        leading=17,
        textColor=colors.HexColor("#0284c7"),
        spaceBefore=12,
        spaceAfter=6
    )
    h2_style = ParagraphStyle(
        'H2',
        parent=styles['Heading3'],
        fontName='Helvetica-Bold',
        fontSize=10.5,
        leading=14,
        textColor=colors.HexColor("#0f172a"),
        spaceBefore=8,
        spaceAfter=4
    )
    body_style = ParagraphStyle(
        'Body',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=13,
        textColor=colors.HexColor("#1e293b"),
        spaceAfter=6
    )
    code_style = ParagraphStyle(
        'CodeStyle',
        parent=styles['Code'],
        fontName='Courier',
        fontSize=6.5,
        leading=8.5,
        textColor=colors.HexColor("#0f172a")
    )
    table_cell = ParagraphStyle(
        'TableCell',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        leading=10,
        textColor=colors.HexColor("#1e293b")
    )
    table_hdr = ParagraphStyle(
        'TableHdr',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8.5,
        leading=11,
        textColor=colors.white
    )

    story = []

    # Title block
    story.append(Paragraph("🏥 MediBot System Architecture & Code Flow", title_style))
    story.append(Paragraph("Enterprise Medical Assistant with Database-Layer RBAC, Hybrid Search & SQL RAG<br/><b>MediAssist Health Network</b>", subtitle_style))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#cbd5e1"), spaceAfter=10))

    # 1. Executive Summary
    story.append(Paragraph("1. Executive Summary", h1_style))
    story.append(Paragraph(
        "<b>MediBot</b> is an enterprise medical and operational AI assistant built for MediAssist Health Network "
        "(12 hospitals and 40+ clinics in India). It fulfills two mission-critical requirements: (1) high-precision retrieval across "
        "medical protocols, drug formularies, nursing standards, equipment manuals, and insurance tariffs; and (2) zero-trust Role-Based "
        "Access Control (RBAC) enforced directly at the vector database retrieval level to eliminate data leakage even under adversarial jailbreaks.",
        body_style
    ))

    # 2. Visual Architecture Diagram
    story.append(Paragraph("2. End-to-End System Flow Diagram", h1_style))
    flowchart_ascii = (
        "+-----------------------------------------------------------------------------------------+\n"
        "|  Staff User (Doctor / Nurse / Billing / Technician / Admin)                              |\n"
        "|  POST /chat { question, role, auth_token }                                               |\n"
        "+--------------------------------------------+--------------------------------------------+\n"
        "                                             |                                             \n"
        "                                             v                                             \n"
        "                       [ FastAPI Backend Gateway (backend/main.py) ]                      \n"
        "                                             |                                             \n"
        "                        +--------------------+--------------------+                        \n"
        "                        | 1. Security Guardrails & Verification  |                        \n"
        "                        | • Gibberish filter & API-key protection|                        \n"
        "                        +--------------------+--------------------+                        \n"
        "                                             |                                             \n"
        "                                 [ Intent Classifier ]                                     \n"
        "                              /                         \\                                  \n"
        "               (Analytical/Claims Question)         (Policy/Clinical Manual Question)     \n"
        "                            /                             \\                                \n"
        "                           v                               v                               \n"
        "     +--------------------------------+           +--------------------------------+       \n"
        "     | BRANCH A: SQL RAG ENGINE       |           | BRANCH B: HYBRID DOCUMENT RAG  |       \n"
        "     | (backend/rag/sql_rag.py)       |           | (backend/rag/hybrid_retriever) |       \n"
        "     +--------------------------------+           +--------------------------------+       \n"
        "     | 1. RBAC Perms Check            |           | 1. Collection RBAC Pre-Check   |       \n"
        "     |    (Billing Exec / Admin only) |           | 2. Qdrant Vector Filter        |       \n"
        "     | 2. NL-to-SQL Generation (LLM)  |           |    must=[access_roles == role] |       \n"
        "     | 3. Safety Sanitizer (SELECT)   |           | 3. Dual Dense + Sparse Search  |       \n"
        "     | 4. Query SQLite mediassist.db  |           | 4. Reciprocal Rank Fusion (RRF)|       \n"
        "     | 5. LLM Answer Synthesis        |           | 5. Cross-Encoder Rerank (Top 3)|       \n"
        "     +----------------+---------------+           | 6. Grounded LLM Synthesis      |       \n"
        "                      |                           +----------------+---------------+       \n"
        "                      \\                                           /                        \n"
        "                       \\-------------------+---------------------/                         \n"
        "                                           v                                               \n"
        "                       [ Structured JSON Response Returned to UI ]                         \n"
        "                       • Grounded Answer text with citations & badges                     \n"
        "                       • Sources: [{ document, section, score }]                          \n"
        "                       • Security Guardrail & RBAC violation flags                         \n"
        "+-----------------------------------------------------------------------------------------+"
    )
    t_code = Table([[Paragraph(f"<pre>{flowchart_ascii}</pre>", code_style)]], colWidths=[532])
    t_code.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor("#f8fafc")),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor("#cbd5e1")),
        ('TOPPADDING', (0,0), (-1,-1), 4),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
        ('LEFTPADDING', (0,0), (-1,-1), 6),
        ('RIGHTPADDING', (0,0), (-1,-1), 6),
    ]))
    story.append(t_code)
    story.append(Spacer(1, 8))

    # 3. Phase 1
    story.append(Paragraph("3. Phase 1: Ingestion Pipeline (Data Preparation)", h1_style))
    story.append(Paragraph(
        "• <b>Document Parsing (backend/ingestion/parser.py):</b> Parses PDFs and Markdown files preserving headings, bullet points, and tables. Attaches hierarchical metadata breadcrumbs like <i>[Document > Section > Subsection]</i>.<br/>"
        "• <b>Role Tagging:</b> Every chunk is tagged with permissible access roles (e.g., clinical guides get <i>doctor, admin</i>; tariffs get <i>billing_executive, admin</i>).<br/>"
        "• <b>Dual Embeddings (backend/ingestion/embedder.py):</b> Dense embeddings (384d, <i>bge-small-en-v1.5</i>) for semantic matching plus Sparse BM25 embeddings (<i>Qdrant/bm25</i>) for medical keywords.<br/>"
        "• <b>Vector Database (backend/ingestion/index_qdrant.py):</b> Upserts points into Qdrant collection <i>mediassist_knowledge</i> with dedicated payload indexes.",
        body_style
    ))

    # 4. Phase 2
    story.append(Paragraph("4. Phase 2: Step-by-Step Runtime Code Flow", h1_style))
    story.append(Paragraph("<b>Step 1: Request Gateway & Pre-Flight Guardrails</b>", h2_style))
    story.append(Paragraph("FastAPI receives the request at <code>/chat</code>. It checks for gibberish patterns (vowel ratios, character repetitions) and blocks malicious prompt injections attempting to extract system API keys.", body_style))

    story.append(Paragraph("<b>Step 2: Intent Classification & Routing</b>", h2_style))
    story.append(Paragraph("Identifies if the query is <i>Analytical / Statistical</i> (contains terms like count, claims, sum, average, tickets) or <i>Document-based</i> (treatment guidelines, dosages, manuals).", body_style))

    story.append(Paragraph("<b>Branch A: SQL RAG Engine (Analytical Queries)</b>", h2_style))
    story.append(Paragraph(
        "1. <b>RBAC Check:</b> Restricted to <code>billing_executive</code> and <code>admin</code>. Unauthorized roles receive a clear refusal.<br/>"
        "2. <b>NL to SQL:</b> LLM converts the request into standard SQLite targeting <code>claims</code> or <code>maintenance_tickets</code>.<br/>"
        "3. <b>SQL Sanitizer:</b> Enforces read-only safety (must start with SELECT; drops DDL/DML like DROP, DELETE, UPDATE).<br/>"
        "4. <b>Execution & Synthesis:</b> Queries <code>mediassist.db</code> and uses the LLM to summarize the returned records in natural language.",
        body_style
    ))

    story.append(Paragraph("<b>Branch B: Hybrid Document RAG (Policies & Manuals)</b>", h2_style))
    story.append(Paragraph(
        "1. <b>Domain Pre-Check:</b> Rejects queries explicitly asking for collections outside the user's scope.<br/>"
        "2. <b>Database-Level RBAC Filter:</b> Enforced inside the Qdrant retrieval query (<code>must=[FieldCondition(key='access_roles', match=MatchValue(value=role))]</code>). Restricted chunks are never retrieved into memory.<br/>"
        "3. <b>Hybrid Search & RRF:</b> Qdrant evaluates dense and sparse BM25 vectors simultaneously and fuses candidate chunks using Reciprocal Rank Fusion (RRF) to retrieve the top 10 candidates.<br/>"
        "4. <b>Cross-Encoder Reranking:</b> <code>cross-encoder/ms-marco-MiniLM-L-6-v2</code> jointly re-scores the query and candidates to pick the top 3 chunks.<br/>"
        "5. <b>LLM Grounded Synthesis:</b> LLM synthesizes an answer referencing only verified citations.",
        body_style
    ))

    # 5. RBAC Matrix Table
    story.append(Paragraph("5. User Roles & Access Control Matrix", h1_style))
    table_data = [
        [Paragraph("Role", table_hdr), Paragraph("Department", table_hdr), Paragraph("Accessible Collections", table_hdr), Paragraph("SQL RAG?", table_hdr), Paragraph("Demo User", table_hdr)],
        [Paragraph("doctor", table_cell), Paragraph("Clinical Cardiology", table_cell), Paragraph("clinical, nursing, general", table_cell), Paragraph("❌ No", table_cell), Paragraph("dr.mehta", table_cell)],
        [Paragraph("nurse", table_cell), Paragraph("ICU & Patient Care", table_cell), Paragraph("nursing, general", table_cell), Paragraph("❌ No", table_cell), Paragraph("nurse.priya", table_cell)],
        [Paragraph("billing_executive", table_cell), Paragraph("Billing & Insurance", table_cell), Paragraph("billing, general", table_cell), Paragraph("✅ Yes", table_cell), Paragraph("billing.ravi", table_cell)],
        [Paragraph("technician", table_cell), Paragraph("Biomedical Eng.", table_cell), Paragraph("equipment, general", table_cell), Paragraph("❌ No", table_cell), Paragraph("tech.anand", table_cell)],
        [Paragraph("admin", table_cell), Paragraph("Executive / IT", table_cell), Paragraph("All 5 collections", table_cell), Paragraph("✅ Yes", table_cell), Paragraph("admin.sys", table_cell)]
    ]
    t_matrix = Table(table_data, colWidths=[80, 110, 172, 70, 100])
    t_matrix.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#0284c7")),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#cbd5e1")),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor("#f8fafc")]),
        ('TOPPADDING', (0,0), (-1,-1), 4),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
    ]))
    story.append(t_matrix)
    story.append(Spacer(1, 8))

    # 6. Core Files Table
    story.append(Paragraph("6. Key Architecture Components & Source Files", h1_style))
    files_table_data = [
        [Paragraph("Module File", table_hdr), Paragraph("Layer", table_hdr), Paragraph("Functionality", table_hdr)],
        [Paragraph("backend/main.py", table_cell), Paragraph("API Gateway", table_cell), Paragraph("FastAPI endpoints (/chat, /login, /collections, /health).", table_cell)],
        [Paragraph("backend/ingestion/parser.py", table_cell), Paragraph("Ingestion", table_cell), Paragraph("Parses PDFs and Markdown, builds hierarchical context & chunks.", table_cell)],
        [Paragraph("backend/ingestion/embedder.py", table_cell), Paragraph("Embedding", table_cell), Paragraph("Generates dense (bge-small-en) & sparse (BM25) vector representations.", table_cell)],
        [Paragraph("backend/ingestion/index_qdrant.py", table_cell), Paragraph("Vector Storage", table_cell), Paragraph("Creates collections, defines schemas, and upserts chunk points.", table_cell)],
        [Paragraph("backend/rag/hybrid_retriever.py", table_cell), Paragraph("Retrieval Core", table_cell), Paragraph("Executes Qdrant RBAC filtering, RRF fusion, and pipeline coordination.", table_cell)],
        [Paragraph("backend/rag/reranker.py", table_cell), Paragraph("Reranking", table_cell), Paragraph("Cross-Encoder reranking using ms-marco-MiniLM-L-6-v2 down to Top 3.", table_cell)],
        [Paragraph("backend/rag/sql_rag.py", table_cell), Paragraph("SQL Analytics", table_cell), Paragraph("NL-to-SQL generation, query sanitizer, SQLite execution on claims/tickets.", table_cell)],
        [Paragraph("backend/rag/llm_service.py", table_cell), Paragraph("LLM Inference", table_cell), Paragraph("Multi-model client supporting Gemini, Groq, OpenAI & local fallback.", table_cell)]
    ]
    t_files = Table(files_table_data, colWidths=[130, 82, 320])
    t_files.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#0f172a")),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#cbd5e1")),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor("#f8fafc")]),
        ('TOPPADDING', (0,0), (-1,-1), 3),
        ('BOTTOMPADDING', (0,0), (-1,-1), 3),
    ]))
    story.append(t_files)

    doc.build(story, canvasmaker=NumberedCanvas)
    print(f"PDF document saved to: {filename}")

if __name__ == "__main__":
    docx_path = OUTPUT_DIR / "MediBot_Architecture_and_Code_Flow.docx"
    pdf_path = OUTPUT_DIR / "MediBot_Architecture_and_Code_Flow.pdf"
    create_docx(docx_path)
    create_pdf(pdf_path)
